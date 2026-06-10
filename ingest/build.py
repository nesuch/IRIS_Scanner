import json, sys, os
from collections import Counter
from .extract import extract
from .preprocess import explode_layout_tables
from .segment import segment
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

def validate(blocks, dropped, inscope, excluded, assigned, rows):
    lines=blocks
    kept={l['n'] for l in lines}
    excl={n for b in excluded for n in b['lines']}
    inscope_ns={l['n'] for l in inscope}
    orphans=sorted(inscope_ns - assigned - excl)
    accounted=assigned | excl | set(orphans)
    # full orphan detail (text + page) so the UI can show exactly what's unaccounted
    _bytext={l['n']: l for l in lines}
    orphan_detail=[]
    for n in orphans:
        l=_bytext.get(n, {})
        orphan_detail.append({'n':n,'page':l.get('page'),
                              'text':(l.get('text') or '[table]')[:300]})
    return {
      'pdf_lines_kept':len(lines),
      'dropped':dict(Counter(d['reason'] for d in dropped)),
      'clauses':len(rows),
      'excluded_blocks':[{'reason':b['reason'],'lines':len(b['lines'])} for b in excluded if b['lines']],
      'orphan_lines':len(orphans),
      'empty_result': (len(rows)==0 and len(inscope)>5),
      'orphan_samples':[next((l.get('text','[table]') for l in lines if l['n']==n),'')[:70] for n in orphans[:6]],
      'orphan_detail':orphan_detail,
      'fully_accounted': accounted==kept,
    }

def _rich(text):
    """Turn $...$ segments into bold runs for an Excel cell. Returns a plain str
    if there are no markers (or on any failure), else a CellRichText."""
    if '$' not in text or text.count('$') < 2:
        return text
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        import re as _re
        parts = _re.split(r'\$(.+?)\$', text)   # odd indices are the bolded pieces
        rt = CellRichText()
        for idx, seg in enumerate(parts):
            if seg == '':
                continue
            if idx % 2 == 1:
                rt.append(TextBlock(InlineFont(b=True), seg))
            else:
                rt.append(seg)
        return rt if len(rt) else text
    except Exception:
        return text.replace('$', '')   # fallback: drop markers, keep words

def to_xlsx(rows, sheet, path):
    wb=Workbook(); ws=wb.active; ws.title=sheet[:31]
    ws.append(['Clause_ID','Clause_Text','Regulatory_Tags'])
    for c in ws[1]: c.font=Font(bold=True)
    for r in rows: ws.append([r['id'], _rich(r['clause']), r['tag']])
    ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=95; ws.column_dimensions['C'].width=40
    for row in ws.iter_rows(min_row=2): row[1].alignment=Alignment(wrap_text=True,vertical='top')
    wb.save(path)

def run(spec_path, pdf_path, out_path):
    spec=json.load(open(spec_path))
    L,D=extract(pdf_path, ignore_patterns=spec.get('ignore_patterns'))
    for step in spec.get('preprocess', []):
        if step == 'explode_layout_tables':
            L = explode_layout_tables(L)
    rows,ins,exc,asg,spec_errors=segment(L,spec)
    rep=validate(L,D,ins,exc,asg,rows)
    os.makedirs(os.path.dirname(out_path),exist_ok=True)
    to_xlsx(rows,spec['sheet_name'],out_path)
    return rows,rep

if __name__=='__main__':
    rows,rep=run(sys.argv[1],sys.argv[2],sys.argv[3])
    print(json.dumps(rep,indent=2,ensure_ascii=False))
