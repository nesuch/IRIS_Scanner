#!/usr/bin/env python3
"""Generate an Excel "access guide": for every Handbook table we ingest, the
IRIS report view and the exact filter path to reach it.

Reads the PHASE_* config lists in handbook_parts_to_iris so the guide stays in
sync with what's actually loaded, and pulls each table's title from the Handbook
workbooks.

Usage:
    python tools/build_access_guide.py "/path/to/Handbook folder"
"""
import argparse
import os
import re
import sys

import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(__file__))
import handbook_parts_to_iris as H

# How each report view cascades, with a placeholder for the table-specific pick.
PATHS = {
    "Insurer": "Insurer-wise View → pick Insurer → Line of Business = «{pick}» → "
               "Class of Biz → Metric → Fin Year → Quarter",
    "Financials": "Statements & Reports → pick Insurer → Statement = «{pick}» → Fin Year → Quarter",
    "State": "State-wise View → pick State → Line of Business = «{pick}» → "
             "Class of Biz → Metric → Fin Year",
    "Industry": "Industry-wise View → pick Sector → Line of Business = «{pick}» → Metric → Fin Year",
    "Channel": "Channel-wise View → pick Channel → Line of Business = «{pick}» → Metric → Fin Year",
    "Ombudsman": "Ombudsman Centres → pick Centre → Statement = «{pick}» → Fin Year",
    "TPA": "TPA Network → pick TPA → Statement = «{pick}» → Fin Year",
}
SEG = "the segment (Fire / Marine / Motor / Health / …)"
# Better LOB hint for configs whose LOB is read from the table (lob index None).
NONE_PICK = {
    "PHASE14_NESTED": "the plan / category (Linked, Non-Linked, ULIP, Traditional, Individual/Group)",
    "PHASE10_CHANNEL_SEGMENT": SEG,
}

# (config attribute, view, index of the LOB/statement label in the tuple or None)
REGISTRY = [
    ("PHASE1", "Insurer", 2),
    ("PHASE2", "Insurer", None),
    ("PHASE2B_TRANSPOSED", "Insurer", 2),
    ("PHASE2B_CLASS", "Insurer", 2),
    ("PHASE14_NESTED", "Insurer", None),
    ("PHASE5_QUARTERLY", "Insurer", 2),
    ("PHASE3_YEAR_SUBMETRIC", "State", 2),
    ("PHASE3_MATRIX", "State", None),
    ("PHASE3_STATE_SIMPLE", "State", 2),
    ("PHASE3_CLASS", "State", 2),
    ("PHASE12_CROSSTAB", "State", 2),
    ("PHASE4", "Financials", 2),
    ("PHASE7_SEGMENTED", "Financials", 2),
    ("PHASE6_REPORTS", "Financials", 2),
    ("PHASE8_TRANSPOSED", "Financials", 2),
    ("PHASE8_PERIODIC", "Financials", 2),
    ("PHASE9_REPORTS_CLASS", "Financials", 2),
    ("PHASE9B_LIFE_HEALTH", "Financials", 2),
    ("PHASE11_MEASURES", "Financials", 2),
    ("PHASE9C_OMBUDSMAN", "Ombudsman", 2),
    ("PHASE10_CHANNEL_MEASURES", "Channel", 2),
    ("PHASE10_CHANNEL_SEGMENT", "Channel", None),
    ("PHASE10_CHANNEL_CLASS", "Channel", 2),
    ("PHASE13_INDUSTRY_ROWS", "Industry", 3),
    ("PHASE13_INDUSTRY_2KEY", "Industry", 3),
]

# One-off tables wired inline in main() (not in a PHASE list).
ONE_OFFS = [
    ("Part III", "77", "TPA", "Network Hospitals"),
    ("Part III", "78", "State", "Network Providers"),
    ("Part III", "66", "Industry", "Claims Development & Aging"),
    ("Part I", "4", "Industry", SEG),
    ("Part II", "55", "Insurer", "Rural / Social Sector Obligations"),
    ("Part II", "43", "Industry", "Policies Issued"),
]

# Summary-sheet tables (separate converter -> Industry / Country views).
SUMMARY_ROWS = [
    ("Summary", "A-C", "Industry", "(sector totals)",
     "Industry-wise View → pick Sector (Life / General / Indian (All)) → Line of Business → Metric → Fin Year"),
    ("Summary", "D-E", "Country", "(international reinsurance)",
     "Country-wise View → pick Country → Line of Business → Metric → Fin Year"),
]


def table_title(parts_dir, part, sheet, _cache={}):
    wb = _cache.get(part)
    if wb is None:
        path = os.path.join(parts_dir, f"{part}.xlsx")
        wb = _cache[part] = openpyxl.load_workbook(path, read_only=True, data_only=True) \
            if os.path.exists(path) else False
    if not wb:
        return ""
    match = next((s for s in wb.sheetnames if s.strip() == sheet), None)
    if not match:
        return ""
    for row in wb[match].iter_rows(values_only=True, max_row=4):
        for c in row:
            if c and re.search(r"\b(TABLE|STATEMENT)\b", str(c).upper()):
                return re.sub(r"\s+", " ", str(c)).strip()
    return ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("parts_dir", help="Folder with 'Part I.xlsx' … 'Part V.xlsx'")
    ap.add_argument("--out", default=os.path.join("knowledge_base", "IRIS_Handbook_Access_Guide.xlsx"))
    args = ap.parse_args()

    rows = []
    seen = set()

    generic = {SEG, *NONE_PICK.values()}

    def add(part, sheet, view, pick):
        key = (part, sheet, view, pick)
        if key in seen:
            return
        seen.add(key)
        title = table_title(args.parts_dir, part, sheet)
        path = (PATHS.get(view, "{pick}").format(pick=pick))
        # Structured fields for the in-app one-click "Open": the report view and the
        # specific Line of Business / Statement to pre-select (blank if generic).
        specific = "" if pick in generic else pick
        if specific == "General":            # matches the engine's canonical LOB
            specific = "General (All Segments)"
        rows.append({
            "Part": part, "Table": sheet, "Handbook Table Title": title,
            "IRIS View": {"Financials": "Statements & Reports"}.get(view, f"{view}-wise"),
            "How to Access (filter path)": path,
            "Dimension": view, "Pick": specific,
        })

    for attr, view, lob_i in REGISTRY:
        for entry in getattr(H, attr, []):
            part, sheet = entry[0], entry[1]
            if lob_i is not None and lob_i < len(entry) and entry[lob_i]:
                pick = entry[lob_i]
            else:
                pick = NONE_PICK.get(attr, SEG)
            add(part, sheet, view, pick)

    for part, sheet, view, pick in ONE_OFFS:
        add(part, sheet, view, pick)

    df = pd.DataFrame(rows)
    df["_n"] = df["Table"].str.extract(r"(\d+)").astype(float)
    df = df.sort_values(["Part", "_n"]).drop(columns="_n")

    summary_df = pd.DataFrame(
        [{"Part": p, "Table": t, "Handbook Table Title": title,
          "IRIS View": f"{v}-wise", "How to Access (filter path)": path,
          "Dimension": v, "Pick": ""}
         for p, t, v, title, path in SUMMARY_ROWS])
    df = pd.concat([df, summary_df], ignore_index=True)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with pd.ExcelWriter(args.out, engine="openpyxl") as xl:
        df.to_excel(xl, index=False, sheet_name="Access Guide")
        ws = xl.sheets["Access Guide"]
        widths = {"A": 9, "B": 7, "C": 64, "D": 22, "E": 90}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        # F (Dimension) and G (Pick) drive the in-app one-click open; hide them so
        # the downloaded spreadsheet stays clean (pandas still reads hidden cols).
        for col in ("F", "G"):
            ws.column_dimensions[col].hidden = True
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="1A237E")
        ws.freeze_panes = "A2"
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)

    print(f"[+] Wrote {len(df)} table mappings -> {args.out}")


if __name__ == "__main__":
    main()
