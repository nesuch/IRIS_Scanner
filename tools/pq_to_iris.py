#!/usr/bin/env python3
"""Ingest a Parliamentary Question reply (.docx) into IRIS.

Renders the document to HTML with mammoth (bold/italic/underline/tables/lists
preserved), extracts plain text + metadata, copies the original .docx for
download, and inserts a row into the pq_documents table.

Run the Flask app once first so the pq_documents table exists, then:
    python tools/pq_to_iris.py "/path/to/PQ.docx" [--tags "dental, claims"]
"""
import argparse
import io
import os
import re
import shutil
import sqlite3
from datetime import datetime

import mammoth
import docx as docxlib
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(HERE, "iris.db")
STATIC_PQ_DIR = os.path.join(HERE, "static", "documents", "pqs")


def _cell_text(cell):
    """Text of one table cell, including any table nested inside it."""
    parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    for t in cell.tables:
        parts.extend(_table_lines(t))
    return " ".join(parts)


def _table_lines(table):
    """A table as one ' | '-joined line per row. A merged cell repeats the same
    _Cell across the span it covers, so consecutive duplicates are collapsed."""
    lines = []
    for row in table.rows:
        cells, prev = [], None
        for cell in row.cells:
            t = _cell_text(cell)
            if t and t != prev:
                cells.append(t)
            prev = t
        if cells:
            lines.append(" | ".join(cells))
    return lines


def _docx_text(d):
    """Plain text of a .docx in document order, tables included.

    Document.paragraphs skips table cells, so a PQ's data tables (the
    insurer-wise figures the reply defers to the note for PAD for) never
    reached body_text and no search could reach them. Walking the body element
    keeps the header ahead of the tables, which _meta() relies on to read the
    first Subject:/date match.
    """
    out = []
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = Paragraph(child, d).text.strip()
            if t:
                out.append(t)
        elif child.tag == qn("w:tbl"):
            out.extend(_table_lines(Table(child, d)))
    return "\n".join(out)


def parse_docx(src, filename=""):
    """Render a PQ .docx (path or raw bytes) to HTML + text + metadata.
    Returns a dict the CLI and the upload endpoint both use."""
    fobj = io.BytesIO(src) if isinstance(src, (bytes, bytearray)) else open(src, "rb")
    try:
        html = mammoth.convert_to_html(fobj, style_map="u => u").value
    finally:
        if not isinstance(src, (bytes, bytearray)):
            fobj.close()
    d = docxlib.Document(io.BytesIO(src) if isinstance(src, (bytes, bytearray)) else src)
    text = _docx_text(d)
    pq_no, house, subject, doc_date, doc_date_iso, title = _meta(text, filename or "")
    return {"html": html, "text": text, "pq_no": pq_no, "house": house,
            "subject": subject, "doc_date": doc_date, "doc_date_iso": doc_date_iso,
            "title": title}


def _text_to_html(text):
    """Wrap extracted plain text into simple paragraph HTML (PDF has no styling
    we can reliably recover, so blank-line-separated blocks become <p> blocks)."""
    import html as _h
    parts = []
    for block in re.split(r"\n\s*\n", text or ""):
        block = block.strip()
        if not block:
            continue
        parts.append("<p>" + _h.escape(block).replace("\n", "<br>") + "</p>")
    return "\n".join(parts)


def parse_pdf(src, filename=""):
    """Extract a PQ .pdf (path or raw bytes) to HTML + text + metadata.

    Returns the same dict shape as parse_docx so the upload endpoints are
    format-agnostic. Bold/table fidelity is lower than .docx (PDF carries no
    semantic styling), but the reply text and metadata are preserved."""
    import pdfplumber
    fobj = io.BytesIO(src) if isinstance(src, (bytes, bytearray)) else open(src, "rb")
    pages = []
    try:
        with pdfplumber.open(fobj) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    finally:
        if not isinstance(src, (bytes, bytearray)):
            fobj.close()
    text = "\n\n".join(t.strip() for t in pages if t.strip())
    html = _text_to_html(text)
    pq_no, house, subject, doc_date, doc_date_iso, title = _meta(text, filename or "")
    return {"html": html, "text": text, "pq_no": pq_no, "house": house,
            "subject": subject, "doc_date": doc_date, "doc_date_iso": doc_date_iso,
            "title": title}


def parse_pq(src, filename=""):
    """Dispatch to the right parser based on the file extension."""
    if (filename or "").lower().endswith(".pdf"):
        return parse_pdf(src, filename=filename)
    return parse_docx(src, filename=filename)


_MONTH_NAMES = ["january", "february", "march", "april", "may", "june",
                "july", "august", "september", "october", "november", "december"]


def _month_num(name):
    """Month number from a full or abbreviated name ('Sept', 'March'), else 0."""
    n = (name or "").lower().rstrip(".")
    if len(n) < 3:
        return 0
    for i, full in enumerate(_MONTH_NAMES, 1):
        if full.startswith(n):
            return i
    return 0


def _iso_date(s):
    """'14th March 2026' -> '2026-03-14'; '' when the date can't be read.

    doc_date keeps the document's own wording for display. This is the sortable
    twin: as raw text those dates sort by leading digit, which puts '3rd July
    2021' above '14th March 2026' and makes "the most recent PQ on this topic"
    unanswerable. ISO text sorts chronologically under a plain ORDER BY.
    """
    m = re.match(r"\s*(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+)\.?\s+(\d{4})\s*$", s or "")
    if not m:
        return ""
    mon = _month_num(m.group(2))
    if not mon:
        return ""
    try:
        return datetime(int(m.group(3)), mon, int(m.group(1))).strftime("%Y-%m-%d")
    except ValueError:      # e.g. 31st February
        return ""


def _meta(text, filename):
    pq_no = (re.search(r"(?:question\s*no\.?\s*|pq[\s_-]*)(\d{2,6})", filename + " " + text, re.I)
             or re.search(r"\b(\d{3,6})\b", filename))
    pq_no = pq_no.group(1) if pq_no else ""
    house = ("Lok Sabha" if re.search(r"\blok\s*sabha\b|\bls\b", filename + " " + text, re.I)
             else "Rajya Sabha" if re.search(r"\brajya\s*sabha\b|\brs\b", filename + " " + text, re.I)
             else "")
    starred = "Starred" if re.search(r"starred", text, re.I) else "Unstarred" if re.search(r"unstarred", text, re.I) else ""
    sm = re.search(r"Subject:\s*(.+)", text)
    subject = re.sub(r"\s+", " ", sm.group(1)).strip().strip('"“”').strip()[:400] if sm else ""
    dm = re.search(r"(\d{1,2}(?:st|nd|rd|th)?\s+[A-Z][a-z]+\s+\d{4})", text)
    doc_date = dm.group(1) if dm else ""
    house_label = f"{house} {starred} Q".strip()
    title = subject or f"{house_label} No. {pq_no}".strip()
    return pq_no, house, subject, doc_date, _iso_date(doc_date), title


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx", help="Path to the PQ .docx file")
    ap.add_argument("--tags", default="", help="Comma-separated tags")
    args = ap.parse_args()

    parsed = parse_docx(args.docx, os.path.basename(args.docx))
    html, text = parsed["html"], parsed["text"]
    pq_no, house, subject, doc_date, doc_date_iso, title = (
        parsed["pq_no"], parsed["house"], parsed["subject"], parsed["doc_date"],
        parsed["doc_date_iso"], parsed["title"])

    os.makedirs(STATIC_PQ_DIR, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", os.path.basename(args.docx))
    shutil.copy(args.docx, os.path.join(STATIC_PQ_DIR, safe))

    conn = sqlite3.connect(DB_PATH)
    conn.execute("""CREATE TABLE IF NOT EXISTS pq_documents (
        id INTEGER PRIMARY KEY AUTOINCREMENT, pq_no TEXT, house TEXT, title TEXT,
        subject TEXT, doc_date TEXT, doc_date_iso TEXT, tags TEXT, html TEXT,
        body_text TEXT, docx_filename TEXT, created_at DATETIME)""")
    conn.execute("""INSERT INTO pq_documents
        (pq_no, house, title, subject, doc_date, doc_date_iso, tags, html, body_text,
         docx_filename, created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (pq_no, house, title, subject, doc_date, doc_date_iso, args.tags.strip(), html,
         text, safe, datetime.utcnow().isoformat(sep=" ", timespec="seconds")))
    conn.commit()
    conn.close()
    print(f"[+] Ingested PQ {pq_no or '?'} ({house}) — {title[:60]!r}")
    print(f"    html={len(html)} chars | tables={html.count('<table>')} | "
          f"strong={html.count('<strong>')} em={html.count('<em>')} u={html.count('<u>')} | docx={safe}")


if __name__ == "__main__":
    main()
