#!/usr/bin/env python3
"""Convert IRDAI Handbook 'Summary' sheets into IRIS tidy financial rows.

The Handbook ships ~100 print-formatted statistical tables. The Summary sheets
A/B/C are clean industry time-series (Particular | Remarks | Unit | <years...>).
This script un-pivots A/B/C into the long format IRIS's data engine ingests and
writes a single tidy .xlsx into knowledge_base/raw_submissions/, so the existing
Admin -> "Sync Data" flow picks it up automatically (no engine changes needed).

Output columns: dimension, entity, metric, value, financial_year, quarter,
                line_of_business, class_of_business

Usage:
    python tools/handbook_to_iris.py "/path/to/Summary.xlsx"
    python tools/handbook_to_iris.py "/path/to/Summary.xlsx" --out knowledge_base/raw_submissions
"""
import argparse
import os
import re
import sys

import openpyxl
import pandas as pd

# Summary sheet -> the industry entity it describes.
SHEET_ENTITY = {
    "A": "Indian Insurance Sector (All)",
    "B": "Life Insurance Sector",
    "C": "General Insurance Sector",
}
YEAR_RE = re.compile(r"^\d{4}-\d{2}$")        # e.g. 2024-25
DIMENSION = "Industry"
QUARTER = "Annual"


def _clean_text(v):
    return re.sub(r"\s+", " ", str(v or "").replace("\n", " ")).strip()


def _clean_metric(particular, unit):
    name = _clean_text(particular).rstrip("#").strip()
    unit = _clean_text(unit)
    # Keep the unit on the metric name (IRIS has no unit column) when meaningful.
    if unit and unit.lower() not in {"none", "nan", "unit", ""}:
        name = f"{name} ({unit})"
    return name


def _to_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        # Strip floating-point noise (e.g. 3.9000000000001 -> 3.9).
        return round(float(v), 4)
    s = str(v).strip().replace(",", "")
    if s in {"", "-", "NA", "N.A.", "na", "nan", "None"}:
        return None
    try:
        return round(float(s), 4)
    except ValueError:
        return None


def _find_header(ws, max_scan=6):
    """Return (header_row_index, {col_index: financial_year}, particular_col, unit_col)."""
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        first = _clean_text(row[0]).lower() if row else ""
        if first.startswith("particular"):
            years, particular_col, unit_col = {}, 0, None
            for ci, cell in enumerate(row):
                txt = _clean_text(cell)
                if YEAR_RE.match(txt):
                    years[ci] = txt
                elif txt.lower() == "unit":
                    unit_col = ci
            if years:
                return r, years, particular_col, unit_col
    return None


def convert_sheet(ws, entity, source):
    info = _find_header(ws)
    if not info:
        print(f"   [!] no header/year row found in '{ws.title}' — skipped")
        return []
    header_row, year_cols, p_col, unit_col = info
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        particular = row[p_col] if p_col < len(row) else None
        if not _clean_text(particular):
            continue
        unit = row[unit_col] if (unit_col is not None and unit_col < len(row)) else None
        metric = _clean_metric(particular, unit)
        for ci, fy in year_cols.items():
            if ci >= len(row):
                continue
            value = _to_number(row[ci])
            if value is None:
                continue  # section headers / footnotes / blanks emit nothing
            rows.append({
                "dimension": DIMENSION,
                "entity": entity,
                "metric": metric,
                "value": value,
                "financial_year": fy,
                "quarter": QUARTER,
                "line_of_business": "General",
                "class_of_business": "General",
            })
    return rows


def main():
    ap = argparse.ArgumentParser(description="Convert IRDAI Handbook Summary into IRIS tidy rows.")
    ap.add_argument("summary_xlsx", help="Path to the Handbook 'Summary.xlsx'")
    ap.add_argument("--out", default=os.path.join("knowledge_base", "raw_submissions"),
                    help="Output directory (default: knowledge_base/raw_submissions)")
    ap.add_argument("--name", default="handbook_2024-25_summary.xlsx", help="Output filename")
    args = ap.parse_args()

    if not os.path.exists(args.summary_xlsx):
        sys.exit(f"Input not found: {args.summary_xlsx}")

    wb = openpyxl.load_workbook(args.summary_xlsx, read_only=True, data_only=True)
    source = os.path.basename(args.summary_xlsx)
    all_rows = []
    for sheet, entity in SHEET_ENTITY.items():
        if sheet not in wb.sheetnames:
            print(f"   [!] sheet '{sheet}' not in workbook — skipped")
            continue
        rows = convert_sheet(wb[sheet], entity, source)
        print(f"   sheet {sheet} -> {entity}: {len(rows)} rows")
        all_rows.extend(rows)

    if not all_rows:
        sys.exit("No rows produced — nothing to write.")

    os.makedirs(args.out, exist_ok=True)
    out_path = os.path.join(args.out, args.name)
    df = pd.DataFrame(all_rows, columns=[
        "dimension", "entity", "metric", "value", "financial_year",
        "quarter", "line_of_business", "class_of_business",
    ])
    df.to_excel(out_path, index=False)
    print(f"\n[+] Wrote {len(df)} tidy rows -> {out_path}")
    print(f"    entities: {sorted(df['entity'].unique())}")
    print(f"    years: {sorted(df['financial_year'].unique())}")
    print(f"    metrics: {df['metric'].nunique()} unique")


if __name__ == "__main__":
    main()
